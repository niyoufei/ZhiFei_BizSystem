from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass(frozen=True)
class UploadedContentReaderDependencies:
    document_cls: Any = None
    read_bytes_from_path: Callable[[Path], bytes] | None = None
    extract_pdf_text: Callable[..., str] | None = None
    extract_dxf_text: Callable[[bytes], str] | None = None
    extract_dwg_text: Callable[[bytes, str], str] | None = None
    extract_image_content: Callable[..., str] | None = None


def read_document_bytes_from_path(
    file_path: Path | None,
    *,
    read_bytes_from_path: Callable[[Path], bytes],
) -> bytes:
    if file_path is None:
        return b""
    return read_bytes_from_path(Path(file_path))


def resolve_document_bytes(
    content: bytes | None,
    file_path: Path | None,
    *,
    read_bytes_from_path: Callable[[Path], bytes],
) -> bytes:
    if content is not None:
        return content
    return read_document_bytes_from_path(file_path, read_bytes_from_path=read_bytes_from_path)


def extract_binary_text_snippet(content: bytes, *, max_chars: int = 4000) -> str:
    decoded = content.decode("utf-8", errors="ignore")
    cleaned = "".join(ch if ch.isprintable() else " " for ch in decoded)
    compact = " ".join(cleaned.split())
    if not compact:
        return ""
    return compact[: max(256, int(max_chars))]


def read_uploaded_file_content(
    content: bytes | None,
    filename: str,
    *,
    material_type: object = None,
    file_path: Path | None = None,
    document_cls: Any = None,
    read_bytes_from_path: Callable[[Path], bytes],
    extract_pdf_text: Callable[..., str],
    extract_dxf_text: Callable[[bytes], str],
    extract_dwg_text: Callable[[bytes, str], str],
    extract_image_content: Callable[..., str],
) -> str:
    name = str(filename or "").lower()
    if name.endswith((".txt", ".md", ".csv")):
        if file_path is not None:
            return Path(file_path).read_text(encoding="utf-8", errors="ignore")
        return bytes(content or b"").decode("utf-8", errors="ignore")
    if name.endswith(".docx"):
        if document_cls is None:
            raise ValueError("DOCX 解析不可用：请安装与当前系统架构兼容的 python-docx/lxml。")
        doc = (
            document_cls(str(file_path))
            if file_path is not None
            else document_cls(io.BytesIO(bytes(content or b"")))
        )
        return "\n".join(p.text for p in doc.paragraphs)
    if name.endswith((".doc", ".docm")):
        snippet = extract_binary_text_snippet(
            resolve_document_bytes(
                content,
                file_path,
                read_bytes_from_path=read_bytes_from_path,
            )
        )
        if snippet:
            return snippet
        return f"[DOC资料] 文件: {filename}（当前环境未启用结构化解析，已纳入文件元信息）"
    if name.endswith(".pdf"):
        return extract_pdf_text(
            content,
            filename,
            material_type=material_type,
            file_path=file_path,
        )
    if name.endswith(".json"):
        if file_path is not None:
            return Path(file_path).read_text(encoding="utf-8", errors="ignore")
        return bytes(content or b"").decode("utf-8", errors="ignore")
    if name.endswith((".xlsx", ".xls", ".xlsm")):
        try:
            import openpyxl

            workbook_source: object = (
                str(file_path) if file_path is not None else io.BytesIO(bytes(content or b""))
            )
            wb = openpyxl.load_workbook(workbook_source, read_only=True, data_only=True)
            parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    parts.append("\t".join(str(cell) if cell is not None else "" for cell in row))
            wb.close()
            return "\n".join(parts)
        except Exception as exc:
            raise ValueError(f"Excel 解析失败: {exc}") from exc
    if name.endswith(".dxf"):
        try:
            return extract_dxf_text(
                resolve_document_bytes(
                    content,
                    file_path,
                    read_bytes_from_path=read_bytes_from_path,
                )
            )
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"DXF 解析失败: {exc}") from exc
    if name.endswith(".dwg"):
        return extract_dwg_text(
            resolve_document_bytes(
                content,
                file_path,
                read_bytes_from_path=read_bytes_from_path,
            ),
            filename,
        )
    if name.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")):
        return extract_image_content(content, filename, file_path=file_path)
    snippet = extract_binary_text_snippet(
        resolve_document_bytes(
            content,
            file_path,
            read_bytes_from_path=read_bytes_from_path,
        ),
        max_chars=2000,
    )
    if snippet:
        return snippet
    raise ValueError(
        "仅支持 .txt、.md、.csv、.doc/.docx/.docm、.pdf、.json、.xlsx/.xls/.xlsm、.dxf/.dwg、图片格式"
    )


def read_uploaded_file_content_with_dependencies(
    content: bytes | None,
    filename: str,
    *,
    material_type: object = None,
    file_path: Path | None = None,
    dependencies: UploadedContentReaderDependencies,
) -> str:
    if dependencies.read_bytes_from_path is None:
        raise ValueError("UploadedContentReaderDependencies.read_bytes_from_path 未配置")
    if dependencies.extract_pdf_text is None:
        raise ValueError("UploadedContentReaderDependencies.extract_pdf_text 未配置")
    if dependencies.extract_dxf_text is None:
        raise ValueError("UploadedContentReaderDependencies.extract_dxf_text 未配置")
    if dependencies.extract_dwg_text is None:
        raise ValueError("UploadedContentReaderDependencies.extract_dwg_text 未配置")
    if dependencies.extract_image_content is None:
        raise ValueError("UploadedContentReaderDependencies.extract_image_content 未配置")
    return read_uploaded_file_content(
        content,
        filename,
        material_type=material_type,
        file_path=file_path,
        document_cls=dependencies.document_cls,
        read_bytes_from_path=dependencies.read_bytes_from_path,
        extract_pdf_text=dependencies.extract_pdf_text,
        extract_dxf_text=dependencies.extract_dxf_text,
        extract_dwg_text=dependencies.extract_dwg_text,
        extract_image_content=dependencies.extract_image_content,
    )


def read_uploaded_file_preview_for_project_name(
    content: bytes | None,
    filename: str,
    *,
    file_path: Path | None = None,
    document_cls: Any = None,
    read_bytes_from_path: Callable[[Path], bytes],
    read_uploaded_file_content_impl: Callable[..., str],
    extract_pdf_text_preview: Callable[..., str],
) -> str:
    """为项目名识别提取轻量预览文本，避免自动创建时深解析整份大文件。"""
    name = str(filename or "").lower()
    if name.endswith((".txt", ".md", ".csv", ".json")):
        if file_path is not None:
            return Path(file_path).read_text(encoding="utf-8", errors="ignore")[:12000]
        return bytes(content or b"").decode("utf-8", errors="ignore")[:12000]
    if name.endswith(".docx"):
        if document_cls is None:
            raise ValueError("DOCX 解析不可用：请安装与当前系统架构兼容的 python-docx/lxml。")
        doc = (
            document_cls(str(file_path))
            if file_path is not None
            else document_cls(io.BytesIO(bytes(content or b"")))
        )
        parts: list[str] = []
        total_chars = 0
        for paragraph in doc.paragraphs:
            text = str(paragraph.text or "").strip()
            if not text:
                continue
            parts.append(text)
            total_chars += len(text)
            if total_chars >= 12000 or len(parts) >= 80:
                break
        return "\n".join(parts)
    if name.endswith((".doc", ".docm")):
        snippet = extract_binary_text_snippet(
            resolve_document_bytes(
                content,
                file_path,
                read_bytes_from_path=read_bytes_from_path,
            ),
            max_chars=12000,
        )
        if snippet:
            return snippet
        return f"[DOC资料] 文件: {filename}（当前环境未启用结构化解析，已纳入文件元信息）"
    if name.endswith(".pdf"):
        return extract_pdf_text_preview(
            content,
            filename,
            material_type="tender_qa",
            max_pages=10,
            max_chars=32000,
            ocr_pages=3,
            stop_when_project_name_found=True,
            file_path=file_path,
        )
    if name.endswith(
        (
            ".xlsx",
            ".xls",
            ".xlsm",
            ".dxf",
            ".dwg",
            ".png",
            ".jpg",
            ".jpeg",
            ".webp",
            ".bmp",
            ".tif",
            ".tiff",
        )
    ):
        return read_uploaded_file_content_impl(content, filename, file_path=file_path)[:12000]
    snippet = extract_binary_text_snippet(
        resolve_document_bytes(
            content,
            file_path,
            read_bytes_from_path=read_bytes_from_path,
        ),
        max_chars=12000,
    )
    if snippet:
        return snippet
    raise ValueError(
        "仅支持 .txt、.md、.csv、.doc/.docx/.docm、.pdf、.json、.xlsx/.xls/.xlsm、.dxf/.dwg、图片格式"
    )


def read_uploaded_file_content_for_parse_mode(
    content: bytes | None,
    filename: str,
    *,
    material_type: object,
    parse_mode: str = "full",
    file_path: Path | None = None,
    normalize_material_type: Callable[..., str],
    read_uploaded_file_content_impl: Callable[..., str],
    extract_pdf_text_preview: Callable[..., str],
    extract_boq_tabular_preview_text: Callable[..., str],
    preview_max_pages_by_type: dict[str, int],
    preview_max_chars_by_type: dict[str, int],
    preview_ocr_pages_by_type: dict[str, int],
    preview_max_sheets_by_type: dict[str, int],
    preview_max_rows_by_type: dict[str, int],
    text_max_sheets_by_type: dict[str, int],
    text_max_rows_by_type: dict[str, int],
) -> str:
    normalized_mode = str(parse_mode or "full").strip().lower()
    normalized_type = normalize_material_type(material_type, filename=filename)
    ext = Path(str(filename or "")).suffix.lower()
    if normalized_mode == "preview":
        max_pages = int(preview_max_pages_by_type.get(normalized_type, 4))
        max_chars = int(preview_max_chars_by_type.get(normalized_type, 16000))
        ocr_pages = int(preview_ocr_pages_by_type.get(normalized_type, 2))
        if ext == ".pdf" and normalized_type in {"tender_qa", "drawing", "boq"}:
            return extract_pdf_text_preview(
                content,
                filename,
                material_type=normalized_type,
                max_pages=max_pages,
                max_chars=max_chars,
                ocr_pages=ocr_pages,
                stop_when_project_name_found=False,
                file_path=file_path,
            )
        if normalized_type == "boq" and ext in {".xlsx", ".xls", ".xlsm", ".csv"}:
            preview_text = extract_boq_tabular_preview_text(
                content,
                filename,
                max_sheets=int(preview_max_sheets_by_type.get("boq", 2)),
                max_rows_per_sheet=int(preview_max_rows_by_type.get("boq", 180)),
                file_path=file_path,
            )
            if preview_text:
                return preview_text[:max_chars]
        return read_uploaded_file_content_impl(
            content,
            filename,
            material_type=material_type,
            file_path=file_path,
        )
    if normalized_type == "boq" and ext in {".xlsx", ".xls", ".xlsm", ".csv"}:
        boq_excerpt = extract_boq_tabular_preview_text(
            content,
            filename,
            max_sheets=int(text_max_sheets_by_type.get("boq", 4)),
            max_rows_per_sheet=int(text_max_rows_by_type.get("boq", 600)),
            file_path=file_path,
        )
        if boq_excerpt:
            return boq_excerpt
    return read_uploaded_file_content_impl(
        content,
        filename,
        material_type=material_type,
        file_path=file_path,
    )
